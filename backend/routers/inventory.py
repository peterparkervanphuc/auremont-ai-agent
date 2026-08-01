import io
import zipfile
from datetime import datetime, timedelta

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, UploadFile, status
from sqlalchemy.orm import Session

from backend.core.config import get_settings
from backend.core.deps import require_role
from backend.core.enums import UserRole
from backend.core.minio_client import ensure_bucket, get_minio_client
from backend.core.mysql_client import get_db
from backend.models.inventory_unit import InventoryUnitImage
from backend.repositories.inventory_unit import (
    add_unit_image,
    create_unit,
    delete_unit,
    get_unit,
    get_unit_by_code,
    list_public_units,
    list_units,
    update_unit,
    upsert_unit_from_bulk_row,
)
from backend.schemas.inventory_unit import (
    InventoryBulkUploadPreview,
    InventoryBulkUploadRow,
    InventoryUnitCreate,
    InventoryUnitFilter,
    InventoryUnitResponse,
    InventoryUnitUpdate,
)

# Kênh 1 — cột bắt buộc trong template Excel (CLAUDE.md §6.6.a bước 1)
REQUIRED_COLUMNS = ["Mã căn", "Dự án", "Toà", "Tầng", "Loại căn", "Diện tích", "Giá", "Trạng thái", "Tên file ảnh"]

public_router = APIRouter(prefix="/inventory", tags=["Inventory (Public)"])
admin_router = APIRouter(
    prefix="/admin/inventory", tags=["Inventory (Admin)"], dependencies=[Depends(require_role(UserRole.ADMIN))]
)


def _stale_after() -> datetime:
    return datetime.utcnow() - timedelta(days=get_settings().inventory_stale_after_days)


# --- Public (Customer Portal / Sale lookup) — CLAUDE.md §6.3.a --------------------------------


@public_router.get("", response_model=list[InventoryUnitResponse])
async def list_public_inventory(filters: InventoryUnitFilter = Depends(), db: Session = Depends(get_db)) -> list[InventoryUnitResponse]:
    return list_public_units(db, filters)


@public_router.get("/{unit_id}", response_model=InventoryUnitResponse)
async def get_public_inventory_unit(unit_id: int, db: Session = Depends(get_db)) -> InventoryUnitResponse:
    unit = get_unit(db, unit_id)
    if unit is None or not unit.is_public:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Unit not found")
    return unit


@public_router.get("/images/{image_id}/url")
async def get_inventory_image_url(image_id: int, db: Session = Depends(get_db)) -> dict[str, str]:
    """Presigned MinIO URL cho 1 ảnh — ảnh gốc không public trực tiếp (CLAUDE.md §6.6)."""
    image = db.query(InventoryUnitImage).filter(InventoryUnitImage.id == image_id).first()
    if image is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Image not found")

    url = get_minio_client().presigned_get_object(get_settings().minio_bucket_inventory, image.image_path)
    return {"url": url}


# --- Admin Kênh 2 — Form UI thêm/sửa 1 căn lẻ (CLAUDE.md §6.6.b) -----------------------------


@admin_router.get("", response_model=list[InventoryUnitResponse])
async def admin_list_inventory(db: Session = Depends(get_db)) -> list[InventoryUnitResponse]:
    return list_units(db)


@admin_router.post("", response_model=InventoryUnitResponse, status_code=status.HTTP_201_CREATED)
async def admin_create_unit(payload: InventoryUnitCreate, db: Session = Depends(get_db)) -> InventoryUnitResponse:
    if get_unit_by_code(db, payload.unit_code) is not None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"Mã căn '{payload.unit_code}' đã tồn tại — xác nhận ghi đè hoặc đổi mã.",
        )
    return create_unit(db, payload)


@admin_router.put("/{unit_id}", response_model=InventoryUnitResponse)
async def admin_update_unit(unit_id: int, payload: InventoryUnitUpdate, db: Session = Depends(get_db)) -> InventoryUnitResponse:
    try:
        return update_unit(db, unit_id, payload)
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(exc)) from exc


@admin_router.delete("/{unit_id}", status_code=status.HTTP_204_NO_CONTENT)
async def admin_delete_unit(unit_id: int, db: Session = Depends(get_db)) -> None:
    try:
        delete_unit(db, unit_id)
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(exc)) from exc


@admin_router.post("/{unit_id}/images", response_model=InventoryUnitResponse)
async def admin_upload_unit_image(unit_id: int, file: UploadFile, db: Session = Depends(get_db)) -> InventoryUnitResponse:
    """Kéo-thả ảnh trực tiếp vào form — gắn thẳng với căn này (CLAUDE.md §6.6.b bước 3)."""
    unit = get_unit(db, unit_id)
    if unit is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Unit not found")

    settings = get_settings()
    ensure_bucket(settings.minio_bucket_inventory)
    object_path = f"{unit.unit_code}/{file.filename}"
    data = await file.read()
    get_minio_client().put_object(
        settings.minio_bucket_inventory, object_path, io.BytesIO(data), length=len(data), content_type=file.content_type
    )
    add_unit_image(db, unit_id, object_path)
    return get_unit(db, unit_id)


# --- Admin Kênh 1 — Upload Excel hàng loạt + .zip ảnh (CLAUDE.md §6.6.a) ----------------------


@admin_router.post("/bulk-upload/preview", response_model=InventoryBulkUploadPreview)
async def admin_bulk_upload_preview(excel_file: UploadFile, images_zip: UploadFile | None = None) -> InventoryBulkUploadPreview:
    """Validate cấu trúc cột + match ảnh, trả về preview trước khi Admin xác nhận lưu (bước 2-4)."""
    rows, _ = _parse_excel(await excel_file.read())
    image_names = _list_zip_image_names(await images_zip.read()) if images_zip is not None else set()

    preview_rows = [
        InventoryBulkUploadRow(unit_code=row["Mã căn"], has_image=_match_image(row["Mã căn"], image_names))
        for row in rows
    ]
    missing = sum(1 for r in preview_rows if not r.has_image)
    return InventoryBulkUploadPreview(rows=preview_rows, total=len(preview_rows), missing_images=missing)


@admin_router.post("/bulk-upload/confirm", response_model=list[InventoryUnitResponse])
async def admin_bulk_upload_confirm(
    excel_file: UploadFile, images_zip: UploadFile | None = None, db: Session = Depends(get_db)
) -> list[InventoryUnitResponse]:
    """Ghi vào inventory_units; ảnh lưu MinIO; căn thiếu ảnh vẫn lưu, đánh dấu 'Chưa có hình ảnh' (bước 5)."""
    rows, project_lookup = _parse_excel(await excel_file.read())
    zip_bytes = await images_zip.read() if images_zip is not None else None
    image_map = _extract_zip_images(zip_bytes) if zip_bytes else {}

    settings = get_settings()
    if image_map:
        ensure_bucket(settings.minio_bucket_inventory)

    saved: list[InventoryUnitResponse] = []
    for row in rows:
        unit_code = row["Mã căn"]
        schema = InventoryUnitCreate(
            unit_code=unit_code,
            project_id=project_lookup.get(row["Dự án"], row["Dự án"]),
            building=row.get("Toà"),
            floor=str(row.get("Tầng")) if row.get("Tầng") is not None else None,
            unit_type=row.get("Loại căn"),
            area_sqm=row.get("Diện tích"),
            price=row.get("Giá"),
            status=row.get("Trạng thái") or "available",
        )
        unit = upsert_unit_from_bulk_row(db, schema)

        image_bytes = _match_image_bytes(unit_code, image_map)
        if image_bytes is not None:
            object_path = f"{unit_code}/{unit_code}.jpg"
            get_minio_client().put_object(
                settings.minio_bucket_inventory, object_path, io.BytesIO(image_bytes), length=len(image_bytes)
            )
            add_unit_image(db, unit.id, object_path)

        saved.append(get_unit(db, unit.id))

    return saved


def _parse_excel(content: bytes) -> tuple[list[dict], dict[str, str]]:
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    sheet = workbook.active
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]

    missing = [col for col in REQUIRED_COLUMNS if col not in header]
    if missing:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"File Excel thiếu cột bắt buộc: {', '.join(missing)}",
        )

    rows = []
    for raw_row in sheet.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in raw_row):
            continue
        rows.append(dict(zip(header, raw_row)))

    return rows, {}


def _list_zip_image_names(zip_bytes: bytes) -> set[str]:
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        return {name for name in zf.namelist() if not name.endswith("/")}


def _extract_zip_images(zip_bytes: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        return {name: zf.read(name) for name in zf.namelist() if not name.endswith("/")}


def _match_image(unit_code: str, image_names: set[str]) -> bool:
    return any(name.rsplit("/", 1)[-1].split(".")[0] == unit_code for name in image_names)


def _match_image_bytes(unit_code: str, image_map: dict[str, bytes]) -> bytes | None:
    for name, data in image_map.items():
        if name.rsplit("/", 1)[-1].split(".")[0] == unit_code:
            return data
    return None
